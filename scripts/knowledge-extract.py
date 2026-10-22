"""Extract a known document format from stdin; emit complete Markdown on stdout.

No paths or URLs from the document are opened. Originals remain in the portable
library. Run in the configured parser environment; this helper installs nothing.
"""

import sys
import tempfile


def line(value=""):
    sys.stdout.write(str(value) + "\n")


def cell(value):
    return str(value if value is not None else "").replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def pdf(source):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        document = PdfReader(source)
    except PdfReadError as error:
        raise ValueError("PDF 结构不完整或已损坏，无法读取页面；原件已完整保留，可上传有效版本后重新解析。") from error
    if document.is_encrypted and not document.decrypt(""):
        raise ValueError("PDF 已加密；原件已保存，需要解密后才能解析正文。")
    for number, page in enumerate(document.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            line("## 第 %s 页\n" % number)
            line(text)
            line()


def docx(source):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(source)
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text:
                style = block.style.name if block.style else ""
                if style.startswith("Heading ") and style[-1:].isdigit():
                    line("#" * min(6, int(style[-1:])) + " " + block.text)
                else:
                    line(block.text)
                line()
        elif isinstance(block, Table):
            for index, row in enumerate(block.rows):
                line("| " + " | ".join(cell(column.text) for column in row.cells) + " |")
                if index == 0:
                    line("| " + " | ".join("---" for _ in row.cells) + " |")
            line()
    for section in document.sections:
        for kind, container in [("页眉", section.header), ("页脚", section.footer)]:
            texts = [paragraph.text for paragraph in container.paragraphs if paragraph.text]
            if texts:
                line("### %s\n" % kind)
                line("\n\n".join(texts))


def xlsx(source):
    from openpyxl import load_workbook

    document = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in document.worksheets:
            line("## " + sheet.title + "\n")
            first = True
            for row in sheet.iter_rows(values_only=True):
                if not any(value is not None for value in row):
                    continue
                line("| " + " | ".join(cell(value) for value in row) + " |")
                if first:
                    line("| " + " | ".join("---" for _ in row) + " |")
                    first = False
            line()
    finally:
        document.close()


def slide_shapes(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            slide_shapes(shape.shapes)
        elif shape.has_text_frame:
            line(shape.text)
            line()
        elif shape.has_table:
            for index, row in enumerate(shape.table.rows):
                line("| " + " | ".join(cell(column.text) for column in row.cells) + " |")
                if index == 0:
                    line("| " + " | ".join("---" for _ in row.cells) + " |")
            line()


def pptx(source):
    from pptx import Presentation

    document = Presentation(source)
    for number, slide in enumerate(document.slides, start=1):
        line("## 幻灯片 %s\n" % number)
        slide_shapes(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                line("### 备注\n")
                line(notes.text)
                line()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    parsers = {"pdf": pdf, "docx": docx, "xlsx": xlsx, "pptx": pptx}
    if len(sys.argv) != 2 or sys.argv[1] not in parsers:
        raise ValueError("不支持的文档格式。")
    # Seekable parsers use an anonymous temporary file, rather than buffering an
    # arbitrarily large upload in RAM. It is removed on close, including errors.
    with tempfile.TemporaryFile() as source:
        while True:
            chunk = sys.stdin.buffer.read(1024 * 1024)
            if not chunk:
                break
            source.write(chunk)
        source.seek(0)
        parsers[sys.argv[1]](source)


if __name__ == "__main__":
    try:
        main()
    except ModuleNotFoundError:
        print("服务器文档解析依赖尚未安装，原件已完整保留。", file=sys.stderr)
        sys.exit(2)
    except Exception as error:
        # Avoid dumping document content or local server paths into API errors.
        message = str(error) if isinstance(error, ValueError) else "文档无法解析，文件可能损坏或格式不受支持；原件已保留。"
        print(message, file=sys.stderr)
        sys.exit(2)
